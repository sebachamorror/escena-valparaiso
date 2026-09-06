#!/usr/bin/env python3
"""
Genera una versión .docx de solo texto de src/app/propuesta/content.json,
para revisión o edición donde no hace falta (o no conviene) adjuntar las
imágenes: cada fotografía, captura, maqueta o logo se reemplaza por una
indicación entre corchetes en su lugar.

Misma fuente de datos que la página web, el PDF y el PPTX
(src/app/propuesta/page.tsx / scripts/export-propuesta-pptx.mjs):
un cambio en content.json actualiza los cuatro.

Uso: python3 scripts/export-propuesta-docx.py
Salida: PDF/Propuesta_Difusion_Digital_QUINTA_ESCENA_2027_texto.docx
"""
import json
import re
from pathlib import Path

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

MUTED = RGBColor(0x6B, 0x6B, 0x74)  # --ink-3

ROOT = Path(__file__).resolve().parent.parent
CONTENT = json.loads((ROOT / "src/app/propuesta/content.json").read_text(encoding="utf-8"))
OUT = ROOT / "PDF/Propuesta_Difusion_Digital_QUINTA_ESCENA_2027_texto.docx"

LINK_RE = re.compile(r"\[([^\]]+)\]\((https?://[^\s)]+)\)")
BOLD_RE = re.compile(r"\*\*(.+?)\*\*")


def add_rich(paragraph, text, italic=False):
    """Agrega texto a un párrafo interpretando **negrita** y [texto](url) -> texto (url)."""
    text = LINK_RE.sub(r"\1 (\2)", text)
    pos = 0
    for m in BOLD_RE.finditer(text):
        if m.start() > pos:
            r = paragraph.add_run(text[pos:m.start()])
            r.italic = italic
        r = paragraph.add_run(m.group(1))
        r.bold = True
        r.italic = italic
        pos = m.end()
    if pos < len(text):
        r = paragraph.add_run(text[pos:])
        r.italic = italic


def add_note(doc, text):
    """Indicación entre corchetes de un elemento visual que no se adjunta (foto, captura, maqueta, logo)."""
    p = doc.add_paragraph()
    r = p.add_run(f"[{text}]")
    r.italic = True
    r.font.size = Pt(10)
    r.font.color.rgb = MUTED
    return p


def render_block(doc, block):
    t = block["type"]
    if t == "p":
        p = doc.add_paragraph()
        add_rich(p, block["text"])
    elif t == "h3":
        doc.add_heading(block["text"], level=3)
    elif t == "quote":
        p = doc.add_paragraph()
        add_rich(p, block["text"], italic=True)
        p.paragraph_format.left_indent = Cm(0.75)
    elif t == "list":
        for item in block["items"]:
            p = doc.add_paragraph(style="List Bullet")
            add_rich(p, item)
    elif t == "flow":
        doc.add_paragraph(" → ".join(block["items"]))
    elif t == "table":
        rows = block["rows"]
        table = doc.add_table(rows=1 + len(rows), cols=2)
        table.style = "Light Grid Accent 1"
        hdr = table.rows[0].cells
        hdr[0].text, hdr[1].text = block["headers"]
        for i, (a, b) in enumerate(rows, start=1):
            cells = table.rows[i].cells
            cells[0].text = re.sub(BOLD_RE, r"\1", a)
            cells[1].text = re.sub(BOLD_RE, r"\1", b)
        doc.add_paragraph()
    elif t == "cards":
        for it in block["items"]:
            p = doc.add_paragraph(style="List Bullet")
            r = p.add_run(f"{it['title']}: ")
            r.bold = True
            p.add_run(it["text"])
    elif t == "shot":
        add_note(doc, f"CAPTURA DE PANTALLA — {block.get('caption') or block.get('alt', '')}")
    elif t == "placeholder":
        add_note(doc, f"ESPACIO PARA FOTOGRAFÍA — {block['text']}")
    elif t == "logos":
        add_note(doc, "LOGOS — " + "; ".join(it["label"] for it in block["items"]))
    elif t == "mockups":
        for it in block["items"]:
            add_note(doc, f"MAQUETA ILUSTRATIVA — {it['label']}")
    elif t == "people":
        for it in block["items"]:
            p = doc.add_paragraph(style="List Bullet")
            r = p.add_run("[FOTO] ")
            r.italic = True
            r2 = p.add_run(it["name"])
            r2.bold = True
            p.add_run(f" — {it['role']}")
        if block.get("caption"):
            cap = doc.add_paragraph()
            r = cap.add_run(block["caption"])
            r.italic = True
            r.font.size = Pt(9)
    elif t == "twoColLists":
        for side in ("left", "right"):
            data = block[side]
            doc.add_heading(data["heading"], level=4)
            for item in data["items"]:
                doc.add_paragraph(item, style="List Bullet")
    else:
        pass


def main():
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    # Portada
    cover = CONTENT["cover"]
    title = doc.add_heading(level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    eyebrow = doc.add_paragraph()
    r = eyebrow.add_run(cover["eyebrow"].upper())
    r.bold = True
    r.font.size = Pt(9)

    for i, line in enumerate(cover["titleLines"]):
        if i > 0:
            title.add_run().add_break()
        title.add_run(line)
    subtitle = doc.add_paragraph(cover["subtitle"])
    subtitle.runs[0].font.size = Pt(13)
    tagline = doc.add_paragraph()
    r = tagline.add_run(cover["tagline"])
    r.italic = True

    meta_table = doc.add_table(rows=len(cover["meta"]), cols=2)
    for i, (k, v) in enumerate(cover["meta"]):
        cells = meta_table.rows[i].cells
        cells[0].text = k
        cells[1].text = v
    doc.add_paragraph()

    concept_map = cover.get("conceptMap")
    if concept_map:
        labels = ", ".join(n["label"] for n in concept_map["nodes"])
        add_note(
            doc,
            "MAPA CONCEPTUAL — diagrama circular con el isotipo de QUINTA ESCENA al centro, "
            f"recibiendo flechas desde: {labels}. {concept_map.get('caption', '')}",
        )

    doc.add_page_break()

    # Páginas
    for page in CONTENT["pages"]:
        kicker_p = doc.add_paragraph()
        r = kicker_p.add_run(f"{page['num']} · {page['kicker'].upper()}")
        r.bold = True
        r.font.size = Pt(9)
        doc.add_heading(page["title"], level=1)
        if page.get("lead"):
            lead_p = doc.add_paragraph()
            add_rich(lead_p, page["lead"])
            lead_p.runs[0].font.size = Pt(12)
        for block in page["blocks"]:
            render_block(doc, block)
        doc.add_page_break()

    OUT.parent.mkdir(exist_ok=True)
    doc.save(OUT)
    print("Generado:", OUT)


if __name__ == "__main__":
    main()
